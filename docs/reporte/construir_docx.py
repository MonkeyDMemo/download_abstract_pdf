"""Genera la memoria técnica en Word a partir de docs/reporte/reporte.md.

El texto fuente es el Markdown: se edita ahí y este script rehace el .docx,
así que el Word nunca queda desactualizado respecto al repositorio.

No es parte del núcleo (que es solo biblioteca estándar): necesita
python-docx, en un entorno aparte.

    python -m venv .venv-reporte
    .venv-reporte/Scripts/python -m pip install python-docx    # Windows
    .venv-reporte/Scripts/python docs/reporte/construir_docx.py

Entiende el Markdown que usa reporte.md y nada más:
- `#` es el título de la portada; los datos de la portada van en un
  comentario `<!-- portada ... -->` con líneas `clave: valor`;
- `##` y `###` son títulos de nivel 1 y 2;
- párrafos, listas con `-` o `1.`, citas con `>`, tablas con `|`;
- **negritas**, *cursivas* y `código` dentro del texto;
- una figura es un comentario `<!-- figura: nombre | pie -->` seguido de su
  bloque mermaid. GitHub dibuja el bloque; aquí se inserta
  `figuras/nombre.png`, que se exporta del mismo bloque.
"""

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

AQUI = Path(__file__).resolve().parent
RAIZ = AQUI.parent.parent

MORADO = RGBColor(0x4B, 0x3F, 0xA6)
GRIS = RGBColor(0x5B, 0x5F, 0x74)
ANCHO_UTIL_CM = 16.0
ALTO_MAX_CM = 12.5


# ------------------------------------------------------------------ estilos

def preparar_estilos(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    for nombre, tam, antes in (("Heading 1", 16, 18), ("Heading 2", 13, 12)):
        estilo = doc.styles[nombre]
        estilo.font.name = "Calibri"
        estilo.font.size = Pt(tam)
        estilo.font.bold = True
        estilo.font.color.rgb = MORADO
        estilo.paragraph_format.space_before = Pt(antes)
        estilo.paragraph_format.space_after = Pt(6)
        estilo.paragraph_format.keep_with_next = True

    pie = doc.styles["Caption"]
    pie.font.size = Pt(9.5)
    pie.font.italic = True
    pie.font.bold = False
    pie.font.color.rgb = GRIS
    pie.paragraph_format.space_after = Pt(12)

    cita = doc.styles["Quote"]
    cita.font.italic = True
    cita.font.color.rgb = RGBColor(0x33, 0x35, 0x48)
    cita.paragraph_format.left_indent = Cm(1)
    cita.paragraph_format.right_indent = Cm(1)


def configurar_pagina(doc):
    for s in doc.sections:
        s.orientation = WD_ORIENT.PORTRAIT
        s.page_width, s.page_height = Cm(21.59), Cm(27.94)   # carta
        s.left_margin = s.right_margin = Cm(2.5)
        s.top_margin = s.bottom_margin = Cm(2.5)


def campo(parrafo, instruccion, relleno=""):
    """Un campo de Word (PAGE, NUMPAGES, TOC…). Word calcula su valor."""
    corrida = parrafo.add_run()
    inicio = OxmlElement("w:fldChar")
    inicio.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruccion
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    texto = OxmlElement("w:t")
    texto.text = relleno
    fin = OxmlElement("w:fldChar")
    fin.set(qn("w:fldCharType"), "end")
    for e in (inicio, instr, sep, texto, fin):
        corrida._r.append(e)


def pie_y_encabezado(doc, titulo_corto):
    s = doc.sections[0]
    s.different_first_page_header_footer = True
    enc = s.header.paragraphs[0]
    enc.text = titulo_corto
    enc.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for r in enc.runs:
        r.font.size = Pt(8.5)
        r.font.color.rgb = GRIS
    pie = s.footer.paragraphs[0]
    pie.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pie.add_run("Página ")
    campo(pie, "PAGE", "1")
    pie.add_run(" de ")
    campo(pie, "NUMPAGES", "1")
    for r in pie.runs:
        r.font.size = Pt(9)
        r.font.color.rgb = GRIS


def actualizar_campos_al_abrir(doc):
    """Que Word llene el índice y los números de página al abrir."""
    ajustes = doc.settings.element
    e = OxmlElement("w:updateFields")
    e.set(qn("w:val"), "true")
    ajustes.append(e)


# ------------------------------------------------------------------- texto

INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\*[^*\s][^*]*?\*)")


def texto_con_formato(parrafo, texto):
    for trozo in INLINE.split(texto):
        if not trozo:
            continue
        if trozo.startswith("**") and trozo.endswith("**"):
            parrafo.add_run(trozo[2:-2]).bold = True
        elif trozo.startswith("`") and trozo.endswith("`"):
            r = parrafo.add_run(trozo[1:-1])
            r.font.name = "Consolas"
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(0x3A, 0x2F, 0x8F)
        elif trozo.startswith("*") and trozo.endswith("*"):
            parrafo.add_run(trozo[1:-1]).italic = True
        else:
            parrafo.add_run(trozo)


def sombrear(celda, color_hex):
    tc_pr = celda._tc.get_or_add_tcPr()
    sombra = OxmlElement("w:shd")
    sombra.set(qn("w:val"), "clear")
    sombra.set(qn("w:color"), "auto")
    sombra.set(qn("w:fill"), color_hex)
    tc_pr.append(sombra)


def tabla(doc, filas):
    celdas = [[c.strip() for c in f.strip().strip("|").split("|")] for f in filas]
    encabezado, cuerpo = celdas[0], [f for f in celdas[2:]]
    t = doc.add_table(rows=1 + len(cuerpo), cols=len(encabezado))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    anchos = anchos_de_columna(encabezado, cuerpo)
    for k, fila in enumerate(t.rows):
        tr_pr = fila._tr.get_or_add_trPr()
        tr_pr.append(OxmlElement("w:cantSplit"))      # una fila no se parte
        if k == 0:
            tr_pr.append(OxmlElement("w:tblHeader"))  # se repite en cada página
    for j, texto in enumerate(encabezado):
        c = t.cell(0, j)
        c.text = ""
        p = c.paragraphs[0]
        texto_con_formato(p, texto)
        for r in p.runs:
            r.bold = True
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            r.font.size = Pt(10)
        sombrear(c, "4B3FA6")
        c.width = Cm(anchos[j])
    for i, fila in enumerate(cuerpo, start=1):
        for j in range(len(encabezado)):
            c = t.cell(i, j)
            c.text = ""
            p = c.paragraphs[0]
            texto_con_formato(p, fila[j] if j < len(fila) else "")
            for r in p.runs:
                r.font.size = Pt(10)
            p.paragraph_format.space_after = Pt(2)
            c.width = Cm(anchos[j])
            if i % 2 == 0:
                sombrear(c, "F3F1FB")
    doc.add_paragraph()


def anchos_de_columna(encabezado, cuerpo):
    """Reparte el ancho útil según cuánto texto lleva cada columna."""
    n = len(encabezado)
    largos = []
    for j in range(n):
        textos = [encabezado[j]] + [f[j] for f in cuerpo if j < len(f)]
        largos.append(max(6, min(60, sum(len(t) for t in textos) / len(textos))))
    total = sum(largos)
    return [max(2.2, ANCHO_UTIL_CM * l / total) for l in largos]


def figura(doc, nombre, pie):
    ruta = AQUI / "figuras" / (nombre + ".png")
    if not ruta.exists():
        raise SystemExit("Falta la figura %s: expórtala de su bloque mermaid." % ruta)
    ancho_px, alto_px = tamano_png(ruta)
    ancho = ANCHO_UTIL_CM
    if alto_px / ancho_px * ancho > ALTO_MAX_CM:
        ancho = ALTO_MAX_CM * ancho_px / alto_px
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(ruta), width=Cm(ancho))
    leyenda = doc.add_paragraph(style="Caption")
    leyenda.alignment = WD_ALIGN_PARAGRAPH.CENTER
    texto_con_formato(leyenda, pie)


def tamano_png(ruta):
    with open(ruta, "rb") as f:
        cabecera = f.read(24)
    return (int.from_bytes(cabecera[16:20], "big"),
            int.from_bytes(cabecera[20:24], "big"))


def portada(doc, titulo, datos):
    for _ in range(5):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    texto_con_formato(p, titulo)
    for r in p.runs:
        r.font.size = Pt(24)
        r.font.bold = True
        r.font.color.rgb = MORADO
    if datos.get("subtitulo"):
        p = doc.add_paragraph(datos["subtitulo"])
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.runs[0].font.size = Pt(15)
        p.runs[0].font.color.rgb = GRIS
    for _ in range(6):
        doc.add_paragraph()
    for clave, prefijo in (("autor", ""), ("asesor", "Asesor: "),
                           ("institucion", ""), ("version", "")):
        if datos.get(clave):
            p = doc.add_paragraph(prefijo + datos[clave])
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.runs[0].font.size = Pt(12)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # «Contenido» no es un título: si lo fuera, saldría dentro del índice.
    p = doc.add_paragraph()
    r = p.add_run("Contenido")
    r.font.size, r.font.bold, r.font.color.rgb = Pt(16), True, MORADO
    p = doc.add_paragraph()
    campo(p, 'TOC \\o "1-2" \\h \\z \\u',
          "Si no aparece el índice: clic derecho aquí y «Actualizar campo».")
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# ------------------------------------------------------------------ lector

def construir(fuente, salida):
    lineas = fuente.read_text(encoding="utf-8").splitlines()
    doc = Document()
    preparar_estilos(doc)
    configurar_pagina(doc)

    titulo = ""
    datos_portada = {}
    i = 0
    # Título y portada
    while i < len(lineas):
        l = lineas[i]
        if l.startswith("# "):
            titulo = l[2:].strip()
        elif l.strip() == "<!-- portada":
            i += 1
            while lineas[i].strip() != "-->":
                clave, _, valor = lineas[i].partition(":")
                datos_portada[clave.strip()] = valor.strip()
                i += 1
            i += 1
            break
        i += 1
    portada(doc, titulo, datos_portada)
    pie_y_encabezado(doc, datos_portada.get("encabezado")
                     or datos_portada.get("subtitulo", ""))
    actualizar_campos_al_abrir(doc)

    parrafo = []
    lista = None   # (estilo, [items])

    def soltar_parrafo():
        if parrafo:
            texto_con_formato(doc.add_paragraph(), " ".join(parrafo))
            parrafo.clear()

    def soltar_lista():
        nonlocal lista
        if lista:
            estilo, items = lista
            for item in items:
                texto_con_formato(doc.add_paragraph(style=estilo), item)
            lista = None

    while i < len(lineas):
        l = lineas[i]
        s = l.strip()
        fig = re.match(r"<!-- figura: (\w+) \| (.*?) -->", s)
        if fig:
            soltar_parrafo(); soltar_lista()
            i += 1
            if i < len(lineas) and lineas[i].strip().startswith("```mermaid"):
                i += 1
                while not lineas[i].strip().startswith("```"):
                    i += 1
            figura(doc, fig.group(1), fig.group(2))
        elif s.startswith("<!--"):
            while "-->" not in lineas[i]:
                i += 1
        elif s.startswith("```"):
            soltar_parrafo(); soltar_lista()
            i += 1
            bloque = []
            while not lineas[i].strip().startswith("```"):
                bloque.append(lineas[i])
                i += 1
            p = doc.add_paragraph()
            r = p.add_run("\n".join(bloque))
            r.font.name = "Consolas"
            r.font.size = Pt(9)
        elif s.startswith("## "):
            soltar_parrafo(); soltar_lista()
            doc.add_heading(s[3:], level=1)
        elif s.startswith("### "):
            soltar_parrafo(); soltar_lista()
            doc.add_heading(s[4:], level=2)
        elif s.startswith("|"):
            soltar_parrafo(); soltar_lista()
            filas = []
            while i < len(lineas) and lineas[i].strip().startswith("|"):
                filas.append(lineas[i])
                i += 1
            tabla(doc, filas)
            continue
        elif s.startswith("> "):
            soltar_parrafo(); soltar_lista()
            cita = []
            while i < len(lineas) and lineas[i].strip().startswith(">"):
                cita.append(lineas[i].strip()[1:].strip())
                i += 1
            texto_con_formato(doc.add_paragraph(style="Quote"), " ".join(cita))
            continue
        elif re.match(r"^(- |\d+\. )", s):
            soltar_parrafo()
            estilo = "List Bullet" if s.startswith("- ") else "List Number"
            if not lista or lista[0] != estilo:
                soltar_lista()
                lista = (estilo, [])
            lista[1].append(re.sub(r"^(- |\d+\. )", "", s))
        elif s == "":
            soltar_parrafo(); soltar_lista()
        elif lista and l.startswith("  "):
            lista[1][-1] += " " + s
        else:
            soltar_lista()
            parrafo.append(s)
        i += 1
    soltar_parrafo(); soltar_lista()

    salida.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(salida))
    return salida


def main():
    ap = argparse.ArgumentParser(description=(__doc__ or "").splitlines()[0])
    ap.add_argument("--fuente", type=Path, default=AQUI / "reporte.md")
    ap.add_argument("--salida", type=Path,
                    default=RAIZ / "salidas" / "reporte" / "Memoria_tecnica_servicio_social.docx")
    args = ap.parse_args()
    print(construir(args.fuente, args.salida))


if __name__ == "__main__":
    main()
