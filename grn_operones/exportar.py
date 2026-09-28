# -*- coding: utf-8 -*-
"""La base curada y sus conflictos, a CSV. Escritura atomica.

Dos archivos y no uno: el operon curado va a `operones_silver.csv` y lo que
hay que mirar a mano va a `operones_conflictos.csv`. Separarlos es lo que hace
que el segundo sea util: una lista de revision mezclada con 3 000 filas
correctas no la revisa nadie.

Cuando hay datos de PGD sale un tercero, `operones_pgd.csv`: la vista de
operones de pseudomonas.com para todos sus operones, una fila por gen, tal
como la dio PGD y antes de curar.

Aparte está el **catálogo maestro** (`filas_catalogo`), que es el archivo de
referencia para el asesor y el laboratorio: una fila por operón curado con
sus genes, su evidencia, el ID y el nombre que le da cada fuente, el promotor
de CDBProm y si está en el catálogo del paso 1. Lo acompañan una hoja de
fuentes y un léame con las cifras que hay que tener presentes al leerlo.

Los CSV son el producto canónico y salen siempre con biblioteca estándar. El
`.xlsx` del catálogo es comodidad encima: se escribe con openpyxl si está
instalado, y si no, se dice y se sigue.

No imprime: recibe un callable `log`.
"""

import csv
import io
import json
import os
import re

from grn_operones.curar import LOCUS, orden_transcripcion, primer_gen

COLUMNAS_SILVER = [
    "clave_genes", "nombre", "locus_tags", "n_genes", "monocistronico",
    "cadena", "nivel_evidencia",
    "n_fuentes", "fuentes", "pmids", "adyacente", "es_alternativa",
    "promotor_cdbprom", "revisar", "curado_en",
]

COLUMNAS_CONFLICTOS = [
    "clave_genes", "locus_tags", "n_genes", "motivo", "nivel_evidencia",
    "n_fuentes", "fuentes", "pmids",
]

# Lo que obliga a mirar a mano, y por que. El orden es el de urgencia: un
# operon cuyos genes no son adyacentes contradice al genoma y probablemente
# es un error de mapeo; una hebra sin verificar solo dice que faltaba el GFF.
#
# Toda marca que `curar` puede poner tiene su texto. Antes solo tenían texto
# cuatro, y en el silver del 27-sep-2026 unas 61 marcas salían como código
# crudo (`comp_con_cita`) en un archivo pensado para leerse a mano.
MOTIVOS = {
    "no_adyacente": "los genes no son consecutivos en el genoma",
    "hebras_distintas": "los genes no comparten hebra: no comparten promotor",
    "genes_sin_resolver": "algún nombre no se pudo llevar a locus tag",
    "hebra_no_verificada": "sin GFF en caché, la hebra no se comprobó",
    "pendiente_revision": ("BioCyc no declara método y trae un artículo: "
                           "falta decidir si es conocido o predicho"),
    "comp_con_cita": ("predicción con cita: la cita suele ser la del método, "
                      "no una demostración del operón"),
    "sin_evidencia": "BioCyc no declara ningún código de evidencia",
    "locus_sufijo_excluido": ("un locus con sufijo de letra (PA0951a) no está "
                              "en el diccionario"),
    "gen_huerfano": "la fuente cita un gen que su propia consulta no devolvió",
    "nombre_ambiguo": ("un nombre corresponde a varios parálogos y no se "
                       "eligió ninguno"),
}


# La vista de operones de pseudomonas.com (`feature/show/?id=…&view=operons`),
# una fila por gen miembro. Las claves son identificadores; los encabezados
# son texto para quien abre el CSV.
COLUMNAS_PGD = [
    "operon_id", "operon", "locus_tag", "gen", "descripcion", "inicio",
    "fin", "hebra", "tipo", "evidencia", "pmid",
]
ENCABEZADOS_PGD = [
    "ID del operón", "Operón", "Locus tag", "Gen", "Descripción (RefSeq)",
    "Inicio", "Fin", "Hebra", "Tipo (RefSeq)", "Evidencia", "PMID",
]

# Como lo dice la página. PseudoCAP escribe la evidencia de cada operón en
# prosa, y esa prosa no viene en la tabla.
EVIDENCIA_PGD = {
    "DOOR": "Computationally-predicted (DOOR)",
    "PseudoCAP": "Literatura (PseudoCAP)",
}


def cargar_anotacion(ruta):
    """{locus_tag: {simbolo, producto, tipo}} de `genes_pao1.tsv`.

    La tabla de operones de PGD no trae la descripción de cada gen, y la
    anotación de PGD está detrás del mismo desafío de Cloudflare que el resto
    del sitio. Se usa la de RefSeq, que ya está en el repositorio, y el
    encabezado lo dice, porque PGD la redacta distinto: `methionyl-tRNA
    synthetase` contra `methionine--tRNA ligase`.
    """
    salida = {}
    with io.open(ruta, encoding="utf-8", newline="") as f:
        for d in csv.DictReader(f, delimiter="\t", quoting=csv.QUOTE_NONE):
            lt = (d.get("locus_tag") or "").strip()
            if lt:
                salida[lt] = dict((c, (d.get(c) or "").strip())
                                  for c in ("simbolo", "producto", "tipo"))
    return salida


def _orden_operon(id_fuente):
    """`operon-10` después de `operon-2`: el orden textual los invierte."""
    m = re.search(r"(\d+)$", id_fuente or "")
    return (int(m.group(1)) if m else float("inf"), id_fuente or "")


def filas_pgd(con, db, anotacion):
    """La vista de operones de PGD, una fila por gen, para todos sus operones.

    Sale del bronce vigente y no de silver, a propósito: es lo que dijo PGD,
    antes de normalizar, deduplicar o mezclar con otras fuentes. Los genes van
    en el orden del archivo, que es el de la página.

    El PMID de un operón de DOOR es el del método (Mao et al. 2009): el parser
    lo saca de `pmid` para que la curación no lo lea como literatura, y aquí
    vuelve a su columna porque la página lo muestra.
    """
    operones = [f for f in db.bronze_vigente(con) if f["fuente"] == "pgd"]
    operones.sort(key=lambda f: _orden_operon(f["id_fuente"]))
    filas = []
    for f in operones:
        crudo = json.loads(f["registro_raw"]) if f["registro_raw"] else {}
        base = crudo.get("source_database") or f["tipo_evidencia"] or ""
        pmid = f["pmid"] or crudo.get("referencia_metodo") or ""
        for g in crudo.get("genes") or []:
            anot = anotacion.get(g["locus_tag"], {})
            filas.append({
                "operon_id": f["id_fuente"],
                "operon": crudo.get("name") or "",
                "locus_tag": g["locus_tag"],
                "gen": g.get("gene_name") or "",
                "descripcion": anot.get("producto", ""),
                "inicio": g.get("start", ""),
                "fin": g.get("end", ""),
                "hebra": g.get("hebra") or f["cadena"] or "",
                "tipo": anot.get("tipo", ""),
                "evidencia": EVIDENCIA_PGD.get(base, base),
                "pmid": pmid,
            })
    return filas


def _fuentes(fila_id, mapa):
    return ";".join(sorted(set(f for f, _ in mapa.get(fila_id, []))))


def filas_silver(con, db, sin_monocistronicos=False):
    """La capa curada, a filas.

    `sin_monocistronicos` **apaga por omision**: la curacion conserva las
    unidades de un gen porque BioCyc registra 3 774 y descartarlas tiraba la
    mayor parte de esa fuente. Filtrar es una decision de quien lee el archivo,
    no de quien lo construye, asi que vive aqui y no en `curar`.
    """
    mapa = db.fuentes_de_silver(con)
    filas = []
    for s in db.silver_de(con):
        if sin_monocistronicos and s["monocistronico"]:
            continue
        filas.append({
            "clave_genes": s["clave_genes"],
            "nombre": s["nombre"] or "",
            "locus_tags": s["locus_tags"],
            "n_genes": s["n_genes"],
            "monocistronico": "si" if s["monocistronico"] else "no",
            "cadena": s["cadena"] or "",
            "nivel_evidencia": s["nivel_evidencia"],
            "n_fuentes": s["n_fuentes"],
            "fuentes": _fuentes(s["id"], mapa),
            "pmids": s["pmids"] or "",
            "adyacente": "si" if s["adyacente"] else "no",
            "es_alternativa": "si" if s["es_alternativa"] else "no",
            "promotor_cdbprom": "si" if s["promotor_cdbprom"] else "no",
            "revisar": s["revisar"] or "",
            "curado_en": s["curado_en"],
        })
    return filas


def filas_conflictos(con, db):
    """Una fila por operon que necesita ojo humano, con el motivo en prosa."""
    mapa = db.fuentes_de_silver(con)
    filas = []
    for s in db.silver_de(con):
        marcas = [m for m in (s["revisar"] or "").split(";") if m]
        if not marcas:
            continue
        filas.append({
            "clave_genes": s["clave_genes"],
            "nombre": s["nombre"] or "",
            "locus_tags": s["locus_tags"],
            "n_genes": s["n_genes"],
            "motivo": "; ".join(MOTIVOS.get(m, m) for m in marcas),
            "nivel_evidencia": s["nivel_evidencia"],
            "n_fuentes": s["n_fuentes"],
            "fuentes": _fuentes(s["id"], mapa),
            "pmids": s["pmids"] or "",
        })
    return filas


# ------------------------------------------------------- catálogo maestro

COLUMNAS_CATALOGO = [
    "clave_genes", "nombre", "genes", "locus_tags", "orden_verificado",
    "n_genes", "hebra", "nivel_evidencia", "evidencia", "pmids", "fuentes",
    "n_fuentes",
    "id_odb", "nombre_odb", "id_biocyc", "id_door", "nombre_door",
    "id_pseudocap", "nombre_pseudocap",
    "promotor_cdbprom", "gen_con_promotor", "score_cdbprom",
    "monocistronico", "adyacente", "es_alternativa",
    "en_catalogo_paso1", "nombre_catalogo_paso1", "revisar", "motivo",
]
ENCABEZADOS_CATALOGO = [
    "Clave (locus tags ordenados)", "Nombre",
    "Genes (símbolo de genes_pao1.tsv)", "Locus tags",
    "Orden de transcripción verificado", "Número de genes", "Hebra",
    "Nivel de evidencia", "Evidencia por fuente", "PMIDs",
    "Fuentes de operones", "Número de fuentes",
    "ID en ODB", "Nombre en ODB", "ID en BioCyc", "ID en PGD (DOOR)",
    "Nombre en PGD (DOOR)", "ID en PGD (PseudoCAP)",
    "Nombre en PGD (PseudoCAP)",
    "Promotor CDBProm", "Gen con promotor", "Score CDBProm",
    "Un solo gen", "Genes adyacentes", "Unidad alternativa",
    "En el catálogo del paso 1", "Nombre en el catálogo del paso 1",
    "Para revisar", "Motivo",
]

COLUMNAS_FUENTES_CATALOGO = [
    "fuente", "aporta", "descripcion", "estado", "fecha_datos",
    "descargado_en", "origen", "sha256", "bytes", "filas_vigentes",
    "unidades",
]
ENCABEZADOS_FUENTES_CATALOGO = [
    "Fuente", "Aporta", "Qué es", "Estado", "Fecha de los datos",
    "Descargado en", "Origen", "Huella (SHA-256)", "Bytes",
    "Filas vigentes en el bronce", "Unidades del catálogo",
]

# Qué aporta cada fuente. Son las cuatro del encargo con dos papeles: tres
# dan operones y CDBProm da promotores, que marcan operones sin crearlos.
# Contarlo como fuente de operones fabricaría una unidad por promotor, o
# haría que un operón con promotor pareciera confirmado por dos fuentes.
FUENTES_CATALOGO = {
    "odb": ("operones",
            "Operon DataBase v4 (operondb.jp): operones conocidos, con el "
            "artículo que los sostiene"),
    "biocyc": ("operones",
               "BioCyc / PseudoCyc: unidades de transcripción, curadas o "
               "predichas por Pathway Tools"),
    "pgd": ("operones",
            "Pseudomonas Genome DB: predicciones de DOOR y operones de "
            "literatura de PseudoCAP"),
    "cdbprom": ("marca promotores",
                "CDBProm (IIMAS, UNAM): promotores predichos (XGBoost). "
                "Marca el operón cuyo primer gen tiene promotor; no crea "
                "operones"),
}


def _si_no(valor):
    return "sí" if valor else "no"


def _crudo(fila):
    """`registro_raw` como dict. Llega como texto JSON desde la base."""
    try:
        d = json.loads(fila["registro_raw"]) if fila["registro_raw"] else {}
    except ValueError:
        return {}
    return d if isinstance(d, dict) else {}


def _unicos(valores):
    vistos, salida = set(), []
    for v in valores:
        if v and v not in vistos:
            vistos.add(v)
            salida.append(v)
    return salida


def claves_paso1(filas, locus_de):
    """{clave_genes: nombre} del catálogo del paso 1.

    `filas` son las de `grn_bronce.operones.leer()`, con el nombre en su
    escritura original. `locus_de` es `Catalogo.locus_tags`, que aplica la
    expansión única del paso 1. Se reciben de fuera para que este módulo no
    importe `grn_bronce`. La clave sale igual que la de silver: los locus
    tags ordenados como cadena, unidos por `|`.
    """
    salida = {}
    for fila in filas:
        nombre = (fila.get("operon") or "").strip()
        locus = locus_de(nombre) if nombre else []
        if locus:
            salida.setdefault("|".join(sorted(locus)), nombre)
    return salida


def filas_catalogo(con, db, anotacion, paso1=None):
    """El catálogo maestro: una fila por operón curado.

    Junta lo que ya existe, sin SQL nuevo:
    - silver, que es la curación;
    - `fuentes_de_silver`, que dice qué registro de cada fuente respalda
      cada operón;
    - el bronce vigente, del que salen el nombre que le da cada fuente, el
      origen dentro de PGD (DOOR o PseudoCAP) y el score de CDBProm.

    PGD se separa por origen porque mezcla dos cosas que no valen lo mismo:
    `mexAB-oprM` es `operon-89` en DOOR, con el nombre sintético
    `mexA-mexB-oprM`, y cuatro IDs de PseudoCAP, uno por artículo.

    `paso1` es `{clave_genes: nombre}` del catálogo del paso 1, o `None` si
    no se pudo leer. En ese caso la columna queda vacía, porque un «no»
    diría algo que no se comprobó.
    """
    vigente, scores = {}, {}
    for f in db.bronze_vigente(con):
        crudo = _crudo(f)
        if f["fuente"] == "cdbprom":
            scores[f["id_fuente"]] = crudo.get("score", "")
        else:
            vigente[(f["fuente"], f["id_fuente"])] = crudo
    mapa = db.fuentes_de_silver(con)

    filas = []
    for s in db.silver_de(con):
        locus = [x for x in (s["locus_tags"] or "").split("|") if x]
        hebra = s["cadena"] or ""
        orden = orden_transcripcion(locus, hebra)
        ids = {"odb": [], "biocyc": [], "DOOR": [], "PseudoCAP": []}
        nombres = {"odb": [], "DOOR": [], "PseudoCAP": []}
        fuentes = set()
        for fuente, idf in mapa.get(s["id"], []):
            fuentes.add(fuente)
            crudo = vigente.get((fuente, idf))
            if fuente == "pgd":
                if crudo is None:
                    # Registro que ya no está en el bronce vigente: hubo una
                    # extracción después de curar. La hoja de fuentes lo
                    # avisa; aquí se deja constancia en vez de adivinar.
                    ids["DOOR"].append("%s (no vigente)" % idf)
                    continue
                grupo = crudo.get("source_database") or "DOOR"
            else:
                grupo = fuente
            if grupo in ids:
                ids[grupo].append(idf)
            if grupo in nombres and crudo:
                nombres[grupo].append(crudo.get("name") or "")
        gen_promotor = primer_gen(locus, hebra) if s["promotor_cdbprom"] else ""
        marcas = [m for m in (s["revisar"] or "").split(";") if m]
        clave = s["clave_genes"]
        filas.append({
            "clave_genes": clave,
            "nombre": s["nombre"] or "",
            "genes": " ".join(
                anotacion.get(l, {}).get("simbolo") or l for l in orden),
            "locus_tags": "|".join(orden),
            # Verificado quiere decir que `orden_transcripcion` pudo ordenar:
            # hay hebra y todos los locus se dejan numerar. Si no, la lista
            # va en el orden de la fuente, que no siempre es el de
            # transcripción.
            "orden_verificado": _si_no(
                hebra in ("+", "-") and all(LOCUS.match(l) for l in locus)),
            "n_genes": s["n_genes"],
            "hebra": hebra,
            "nivel_evidencia": s["nivel_evidencia"],
            "evidencia": "; ".join(
                c for c in (s["evidencia_codigos"] or "").split(";") if c),
            "pmids": s["pmids"] or "",
            "fuentes": ";".join(sorted(fuentes)),
            "n_fuentes": s["n_fuentes"],
            "id_odb": ";".join(sorted(ids["odb"], key=_orden_operon)),
            "nombre_odb": "; ".join(_unicos(nombres["odb"])),
            "id_biocyc": ";".join(sorted(ids["biocyc"])),
            "id_door": ";".join(sorted(ids["DOOR"], key=_orden_operon)),
            "nombre_door": "; ".join(_unicos(nombres["DOOR"])),
            "id_pseudocap": ";".join(sorted(ids["PseudoCAP"],
                                            key=_orden_operon)),
            "nombre_pseudocap": "; ".join(_unicos(nombres["PseudoCAP"])),
            "promotor_cdbprom": _si_no(s["promotor_cdbprom"]),
            "gen_con_promotor": gen_promotor or "",
            "score_cdbprom": scores.get(gen_promotor, "") if gen_promotor
                             else "",
            "monocistronico": _si_no(s["monocistronico"]),
            "adyacente": _si_no(s["adyacente"]),
            "es_alternativa": _si_no(s["es_alternativa"]),
            "en_catalogo_paso1": ("" if paso1 is None
                                  else _si_no(clave in paso1)),
            "nombre_catalogo_paso1": (paso1 or {}).get(clave, ""),
            "revisar": ";".join(marcas),
            "motivo": "; ".join(MOTIVOS.get(m, m) for m in marcas),
        })
    return filas


def filas_fuentes_catalogo(con, db, fuentes, url_pgd):
    """La hoja de fuentes: de dónde sale cada foto. Una fila por archivo.

    Publica el origen (URL o `archivo-local:<nombre>`), la huella y la fecha
    de descarga, **nunca la ruta en disco**, que es de la máquina de quien
    corrió y no le dice nada al asesor.

    La fecha de descarga no es la de los datos. La tabla de PGD se bajó hoy,
    pero es una foto del 19 de julio de 2021, y eso se dice cuando el origen
    es la tabla fijada (`url_pgd`).
    """
    silver = db.silver_de(con)
    curado_en = max([s["curado_en"] for s in silver] or [""])
    respaldos = {}
    for pares in db.fuentes_de_silver(con).values():
        for fuente in set(f for f, _ in pares):
            respaldos[fuente] = respaldos.get(fuente, 0) + 1
    con_promotor = sum(1 for s in silver if s["promotor_cdbprom"])
    vigentes = {}
    for f in db.bronze_vigente(con):
        vigentes[f["fuente"]] = vigentes.get(f["fuente"], 0) + 1
    ultima = db.ultima_extraccion_completa(con)
    sin_foto = set(db.fuentes_sin_foto(con))

    filas = []
    for fuente in fuentes:
        aporta, descripcion = FUENTES_CATALOGO.get(fuente, ("", ""))
        unidades = (con_promotor if fuente == "cdbprom"
                    else respaldos.get(fuente, 0))
        base = {"fuente": fuente, "aporta": aporta,
                "descripcion": descripcion,
                "filas_vigentes": vigentes.get(fuente, 0),
                "unidades": unidades}
        eid = ultima.get(fuente)
        if not eid:
            base["estado"] = ("fuera de la curación: ninguna extracción "
                              "terminó" if fuente in sin_foto
                              else "sin datos")
            filas.append(base)
            continue
        for d in db.descargas_de_extraccion(con, eid):
            fila = dict(base)
            fila.update({
                "estado": ("más nueva que la curación: corre `curar`"
                           if curado_en and d["descargado_en"] > curado_en
                           else "vigente"),
                "fecha_datos": ("2021-07-19 (exportación del curador de PGD)"
                                if d["url"] == url_pgd else ""),
                "descargado_en": d["descargado_en"],
                "origen": d["url"],
                "sha256": d["sha256"],
                "bytes": d["bytes"],
            })
            filas.append(fila)
    return filas


def lineas_leame(catalogo, fuentes, paso1=None):
    """Lo que hay que saber antes de leer el catálogo, con cifras de hoy.

    Las cifras se calculan al correr y no se escriben a mano: un léame con
    números fijos se queda mintiendo en cuanto cambia una fuente.
    """
    n = len(catalogo)
    if not n:
        return ["El catálogo está vacío: corre `extraer` y `curar` primero."]

    def cuenta(pred):
        return sum(1 for f in catalogo if pred(f))

    predicho = cuenta(lambda f: f["nivel_evidencia"] == "predicho")
    conocido = cuenta(lambda f: f["nivel_evidencia"] == "conocido")
    curado = cuenta(lambda f: f["nivel_evidencia"] == "curado")
    multi = cuenta(lambda f: f["monocistronico"] == "no")
    sin_nombre = cuenta(lambda f: not f["nombre"])
    solo_door = cuenta(lambda f: f["nombre"] and f["nombre_door"]
                       and not f["nombre_odb"] and not f["nombre_pseudocap"])
    con_promotor = cuenta(lambda f: f["promotor_cdbprom"] == "sí")
    varios_ids = cuenta(lambda f: any(
        ";" in f[c] for c in ("id_odb", "id_biocyc", "id_door",
                              "id_pseudocap")))
    sin_orden = cuenta(lambda f: f["orden_verificado"] == "no")
    pct = 100.0 * predicho / n
    pgd_2021 = any(f.get("fecha_datos", "").startswith("2021")
                   for f in fuentes if f["fuente"] == "pgd")

    lineas = [
        "CATÁLOGO MAESTRO DE OPERONES DE P. aeruginosa PAO1",
        "",
        "Una fila por operón curado: %d unidades, %d de más de un gen." % (
            n, multi),
        "",
        "Cuatro fuentes con dos papeles. ODB, BioCyc y PGD aportan operones. "
        "CDBProm aporta promotores predichos: marca %d operones cuyo primer "
        "gen transcrito tiene promotor, y no crea ninguno." % con_promotor,
        "",
        "Cómo leerlo:",
        "- Nivel de evidencia: %d conocidos (con artículo), %d curados y %d "
        "predichos (%.1f %%). La mayor parte es predicción." % (
            conocido, curado, predicho, pct),
        "- «Número de fuentes» cuenta fuentes distintas, no confirmaciones "
        "independientes. PGD (DOOR) y BioCyc (Pathway Tools) son predictores "
        "distintos pero parten de la misma distancia intergénica, y ODB, "
        "BioCyc y PseudoCAP pueden citar el mismo artículo.",
        "- Nombres: %d unidades no tienen nombre (BioCyc no nombra), y %d "
        "solo tienen el de DOOR, que une los nombres de cada gen "
        "(PA0006-lptA) y no es un nombre de la literatura." % (
            sin_nombre, solo_door),
        "- Varios IDs de una misma fuente en una fila (%d filas) son "
        "unidades distintas de esa fuente con los mismos genes: la clave del "
        "catálogo son los genes y ahí colapsan. En PseudoCAP suele ser un ID "
        "por artículo." % varios_ids,
        "- Genes y locus tags van en orden de transcripción cuando se conoce "
        "la hebra. %d filas no la tienen y van en el orden de la fuente." % (
            sin_orden),
    ]
    if pgd_2021:
        lineas.append(
            "- PGD es una foto del 19 de julio de 2021: la exportación de su "
            "curador, publicada por el laboratorio Greene. pseudomonas.com "
            "no ofrece una descarga de operones.")
    if paso1 is not None:
        en_paso1 = cuenta(lambda f: f["en_catalogo_paso1"] == "sí")
        lineas.append(
            "- «En el catálogo del paso 1» compara con operones_pao1.tsv "
            "(%d operones derivados por adyacencia en el genoma, que es otra "
            "predicción). %d unidades coinciden exactamente en genes. Un «sí» "
            "no es evidencia; el paso 1 sigue usando ese catálogo." % (
                len(paso1), en_paso1))
    lineas += [
        "",
        "La hoja de fuentes dice de dónde sale cada foto, con su huella.",
    ]
    return lineas


def escribir_texto(ruta, lineas, log=lambda m: None):
    """Un texto plano, con escritura atómica. El léame canónico."""
    tmp = ruta + ".tmp"
    d = os.path.dirname(ruta)
    if d:
        os.makedirs(d, exist_ok=True)
    with io.open(tmp, "w", encoding="utf-8", newline="\n") as f:
        f.write("\n".join(lineas) + "\n")
    os.replace(tmp, ruta)
    log("  %s" % ruta)


def escribir_csv(ruta, columnas, filas, log=lambda m: None, encabezados=None,
                 bom=False):
    """`encabezados` reemplaza la fila de encabezado sin tocar las claves.

    Así el CSV puede decir «Descripción (RefSeq)» mientras las claves de cada
    fila siguen siendo identificadores ASCII.

    `bom=True` escribe `utf-8-sig`: Excel en Windows abre con doble clic un
    UTF-8 sin BOM como si fuera cp1252, y «Operón» sale «OperÃ³n». Solo se usa
    en los CSV pensados para abrirse en Excel.
    """
    tmp = ruta + ".tmp"
    d = os.path.dirname(ruta)
    if d:
        os.makedirs(d, exist_ok=True)
    with io.open(tmp, "w", encoding="utf-8-sig" if bom else "utf-8",
                 newline="") as f:
        w = csv.DictWriter(f, fieldnames=columnas, extrasaction="ignore")
        if encabezados:
            csv.writer(f).writerow(encabezados)
        else:
            w.writeheader()
        for fila in filas:
            w.writerow(fila)
    os.replace(tmp, ruta)
    log("  %s  (%d filas)" % (ruta, len(filas)))
    return len(filas)


# Anchos a ojo, por NOMBRE de columna y no por letra, por la misma razón que
# en `grn_bronce/exportar.py`: insertar una columna corre las letras y los
# anchos acaban adornando la columna equivocada.
ANCHOS_CATALOGO = {
    "clave_genes": 28, "nombre": 24, "genes": 30, "locus_tags": 28,
    "evidencia": 34, "pmids": 18, "id_biocyc": 16, "id_door": 14,
    "nombre_door": 26, "id_pseudocap": 22, "nombre_pseudocap": 24,
    "nombre_odb": 16, "nombre_catalogo_paso1": 22, "revisar": 24,
    "motivo": 60, "descripcion": 60, "estado": 30, "fecha_datos": 26,
    "descargado_en": 24, "origen": 70, "sha256": 66,
}


def escribir_xlsx_catalogo(ruta, catalogo, fuentes, leame,
                           log=lambda m: None):
    """El catálogo en un libro de tres hojas. Devuelve True si se escribió.

    Si openpyxl no está, se dice y se sigue: los CSV y el léame ya salieron y
    son el producto canónico. El primer import es el del paquete raíz, porque
    es el único que falla limpio cuando openpyxl no está; un submódulo que ya
    estaba en caché importaría aunque el paquete no.

    Tres cosas que openpyxl no perdona o interpreta a su manera:
    - Un carácter de control en una celda lanza `IllegalCharacterError`
      después de que el CSV ya se escribió. Se limpian antes.
    - Un texto que empieza con `=` se guarda como fórmula. Se fuerza a texto.
    - Con el archivo abierto en Excel, Windows no deja reemplazarlo. Se dice
      y se devuelve False en vez de reventar la corrida.
    """
    try:
        from openpyxl import Workbook
    except ImportError:
        log("  openpyxl no está instalado: se omite el .xlsx. Los CSV y el "
            "léame son el producto canónico y ya están escritos.")
        return False
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    def limpio(valor):
        if isinstance(valor, str):
            return ILLEGAL_CHARACTERS_RE.sub("", valor)
        return valor

    def numero(valor):
        try:
            return float(valor)
        except (TypeError, ValueError):
            return valor

    def hoja(libro, titulo, columnas, encabezados, filas, primera=False):
        h = libro.active if primera else libro.create_sheet()
        h.title = titulo
        h.append(encabezados)
        for fila in filas:
            valores = [limpio(fila.get(c, "")) for c in columnas]
            if "score_cdbprom" in columnas:
                i = columnas.index("score_cdbprom")
                valores[i] = numero(valores[i])
            h.append(valores)
            for celda in h[h.max_row]:
                if isinstance(celda.value, str) and celda.value.startswith("="):
                    celda.data_type = "s"
        for celda in h[1]:
            celda.font = Font(bold=True)
        h.freeze_panes = "A2"
        for i, nombre in enumerate(columnas, 1):
            h.column_dimensions[get_column_letter(i)].width = \
                ANCHOS_CATALOGO.get(nombre, 14)
        return h

    libro = Workbook()
    hoja(libro, "Catálogo", COLUMNAS_CATALOGO, ENCABEZADOS_CATALOGO,
         catalogo, primera=True)
    hoja(libro, "Fuentes", COLUMNAS_FUENTES_CATALOGO,
         ENCABEZADOS_FUENTES_CATALOGO, fuentes)
    leame_h = libro.create_sheet("Léame")
    for linea in leame:
        leame_h.append([limpio(linea)])
    leame_h.column_dimensions["A"].width = 120

    d = os.path.dirname(ruta)
    if d:
        os.makedirs(d, exist_ok=True)
    tmp = ruta + ".tmp"
    libro.save(tmp)
    try:
        os.replace(tmp, ruta)
    except OSError as e:
        # PermissionError en Windows cuando Excel tiene el archivo abierto.
        log("  no se pudo reemplazar %s (%s). ¿Está abierto en Excel? Los "
            "CSV y el léame sí se escribieron." % (ruta, e))
        try:
            os.remove(tmp)
        except OSError:
            pass
        return False
    log("  %s  (%d operones)" % (ruta, len(catalogo)))
    return True
